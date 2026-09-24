# /// script
# requires-python = ">=3.10"
# dependencies = ["playwright", "python-pptx"]
# ///
"""Export the HTML decks to PDF, per-slide PNGs, or an image-per-slide PPTX.

    uv run export.py                    # every numbered deck -> out/pdf/<deck>.pdf
    uv run export.py 02_the_cone.html   # just one deck
    uv run export.py --png              # also out/png/<deck>/NN.png (2x resolution)
    uv run export.py --pptx             # also out/pptx/<deck>.pptx (for PowerPoint-only venues)

The HTML files stay the source of truth; everything here is a derived copy.
The first run downloads a headless Chromium through Playwright (~150 MB).
"""

import argparse
import subprocess
import sys
from pathlib import Path

from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import sync_playwright

HERE = Path(__file__).resolve().parent
OUT = HERE / "out"
W, H = 1280, 720


def launch(p):
    try:
        return p.chromium.launch()
    except PlaywrightError:
        print("Installing headless Chromium (one time)...", file=sys.stderr)
        subprocess.run([sys.executable, "-m", "playwright", "install", "chromium"], check=True)
        return p.chromium.launch()


def open_deck(browser, deck: Path, scale: int):
    page = browser.new_page(viewport={"width": W, "height": H}, device_scale_factor=scale)
    page.goto(deck.as_uri() + "?mode=print")
    page.wait_for_function("document.documentElement.dataset.ready === '1'", timeout=30_000)
    return page


def export(deck: Path, browser, png: bool, pptx: bool) -> None:
    name = deck.stem
    page = open_deck(browser, deck, scale=2 if (png or pptx) else 1)
    (OUT / "pdf").mkdir(parents=True, exist_ok=True)
    pdf = OUT / "pdf" / f"{name}.pdf"
    page.pdf(path=str(pdf), width=f"{W}px", height=f"{H}px", print_background=True,
             prefer_css_page_size=True)
    print(f"wrote {pdf.relative_to(HERE)}")
    if not (png or pptx):
        page.close()
        return

    shots = []
    slides = page.locator("section.slide")
    folder = OUT / "png" / name
    folder.mkdir(parents=True, exist_ok=True)
    for i in range(slides.count()):
        path = folder / f"{i + 1:02d}.png"
        slides.nth(i).screenshot(path=str(path))
        notes = slides.nth(i).evaluate(
            "s => Array.from(s.querySelectorAll('aside.notes')).map(n => n.innerText).join('\\n')")
        shots.append((path, notes))
    print(f"wrote {len(shots)} slides to {folder.relative_to(HERE)}/")
    page.close()

    if pptx:
        from pptx import Presentation
        from pptx.util import Emu

        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 13.333 x 7.5 in
        blank = prs.slide_layouts[6]
        for path, notes in shots:
            s = prs.slides.add_slide(blank)
            s.shapes.add_picture(str(path), 0, 0, prs.slide_width, prs.slide_height)
            if notes:
                s.notes_slide.notes_text_frame.text = notes
        (OUT / "pptx").mkdir(parents=True, exist_ok=True)
        out = OUT / "pptx" / f"{name}.pptx"
        prs.save(out)
        print(f"wrote {out.relative_to(HERE)} (images only; edit the HTML, not this)")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("decks", nargs="*", help="deck files (default: every NN_*.html here)")
    ap.add_argument("--png", action="store_true", help="also write one PNG per slide")
    ap.add_argument("--pptx", action="store_true", help="also write an image-per-slide .pptx")
    args = ap.parse_args()

    decks = [Path(d).resolve() for d in args.decks] or sorted(HERE.glob("[0-9][0-9]_*.html"))
    if not decks:
        sys.exit("no decks found")
    with sync_playwright() as p:
        browser = launch(p)
        for deck in decks:
            export(deck, browser, args.png, args.pptx)
        browser.close()


if __name__ == "__main__":
    main()
